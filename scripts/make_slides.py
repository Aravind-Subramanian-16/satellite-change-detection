"""Generate the 12-slide talk deck from the saved paper evidence.

Reads only degradation_results.csv and benchmarks/results/oscd_results_v2.json
(plus the derived oscd_evidence.json). It does not read imagery or score a city.
"""

from __future__ import annotations

import csv
import json
from collections import defaultdict
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "presentation"
ASSETS = OUT / "assets"
CSV_PATH = ROOT / "degradation_results.csv"
OSCD_PATH = ROOT / "benchmarks" / "results" / "oscd_results_v2.json"
EVIDENCE_PATH = ROOT / "benchmarks" / "results" / "oscd_evidence.json"

FONT = "Arial"
NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = RGBColor(0x2E, 0x75, 0xB6)
LIGHT_BLUE = RGBColor(0xEA, 0xF1, 0xF8)
PALE_BLUE = RGBColor(0xD9, 0xE2, 0xF3)
ORANGE = RGBColor(0xC5, 0x5A, 0x11)
LIGHT_ORANGE = RGBColor(0xFB, 0xE5, 0xD6)
TEAL = RGBColor(0x2F, 0x7D, 0x7D)
LIGHT_TEAL = RGBColor(0xE2, 0xF0, 0xEF)
INK = RGBColor(0x26, 0x26, 0x26)
GREY = RGBColor(0x59, 0x59, 0x59)
MID_GREY = RGBColor(0x7F, 0x7F, 0x7F)
LINE = RGBColor(0xD9, 0xE2, 0xF3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FACTORS = [1, 2, 4, 8, 16, 20]
RESOLUTIONS = [0.5, 1.0, 2.0, 4.0, 8.0, 10.0]


def load_degradation() -> dict[int, dict[str, float]]:
    rows = list(csv.DictReader(CSV_PATH.open()))
    by_factor: dict[int, list[dict[str, str]]] = defaultdict(list)
    for row in rows:
        by_factor[int(row["factor"])].append(row)
    pooled: dict[int, dict[str, float]] = {}
    for factor in FACTORS:
        group = by_factor[factor]
        tp = sum(float(r["tp"]) for r in group)
        fp = sum(float(r["fp"]) for r in group)
        fn = sum(float(r["fn"]) for r in group)
        precision = tp / (tp + fp) if tp + fp else 0.0
        recall = tp / (tp + fn) if tp + fn else 0.0
        f1 = 2 * precision * recall / (precision + recall) if precision + recall else 0.0
        iou = tp / (tp + fp + fn) if tp + fp + fn else 0.0
        pooled[factor] = {
            "precision": precision,
            "recall": recall,
            "f1": f1,
            "iou": iou,
            "tp": tp,
            "fp": fp,
            "fn": fn,
            "predicted": tp + fp,
            "changed": tp + fn,
            "images": len(group),
        }
    checks = {
        1: (87.06, 88.82, 85.36),
        2: (70.81, 93.43, 57.01),
        4: (4.41, 90.78, 2.26),
        8: (2.84, 97.69, 1.44),
        16: (2.61, 99.59, 1.32),
        20: (2.61, 99.47, 1.32),
    }
    for factor, (f1, precision, recall) in checks.items():
        assert round(pooled[factor]["f1"] * 100, 2) == f1, factor
        assert round(pooled[factor]["precision"] * 100, 2) == precision, factor
        assert round(pooled[factor]["recall"] * 100, 2) == recall, factor
    assert pooled[20]["images"] == 2048
    assert int(pooled[20]["predicted"]) == 91006
    assert int(pooled[20]["changed"]) == 6837404
    return pooled


def load_oscd() -> tuple[dict, dict]:
    evidence = json.loads(EVIDENCE_PATH.read_text())
    rows = json.loads(OSCD_PATH.read_text())
    assert len(rows) == 50
    assert len({r["city"] for r in rows}) == 10
    expected = {
        "hardcoded": 0.0782,
        "percentile": 0.1176,
        "otsu": 0.0953,
        "cnn_joint": 0.0007,
        "cnn_reflectance": 0.0,
    }
    for method, value in expected.items():
        assert round(evidence["pooled"][method]["f1"], 4) == value, method
    return evidence, rows


def make_f1_chart(pooled: dict[int, dict[str, float]]) -> Path:
    path = ASSETS / "f1_vs_resolution_slide.png"
    fig, ax = plt.subplots(figsize=(9.4, 5.1), dpi=200)
    f1 = [pooled[f]["f1"] * 100 for f in FACTORS]
    recall = [pooled[f]["recall"] * 100 for f in FACTORS]
    ax.plot(RESOLUTIONS, f1, color="#2E75B6", marker="o", markersize=10,
            linewidth=3.2, label="F1")
    ax.plot(RESOLUTIONS, recall, color="#C55A11", marker="s", markersize=8,
            linewidth=2.4, linestyle="--", label="Recall")
    ax.set_xscale("log")
    ax.set_xlim(0.42, 12.5)
    ax.set_xticks(RESOLUTIONS)
    ax.set_xticklabels(["0.5", "1", "2", "4", "8", "10"])
    ax.minorticks_off()
    ax.set_xlabel("Simulated resolution (m)", fontsize=15)
    ax.set_ylabel("Score (%)", fontsize=15)
    ax.set_ylim(0, 100)
    ax.tick_params(axis="both", labelsize=13)
    ax.grid(axis="y", color="#E6E6E6", linewidth=1)
    ax.set_axisbelow(True)
    for spine in ("top", "right"):
        ax.spines[spine].set_visible(False)
    label_points = {0.5, 1.0, 2.0}
    for x, y in zip(RESOLUTIONS, f1):
        if x in label_points:
            if x == 2.0:
                ax.annotate(f"{y:.2f}", (x, y), textcoords="offset points",
                            xytext=(16, 4), ha="left", fontsize=12,
                            color="#2E75B6", fontweight="bold")
            else:
                ax.annotate(f"{y:.2f}", (x, y), textcoords="offset points",
                            xytext=(0, 14), ha="center", fontsize=12,
                            color="#2E75B6", fontweight="bold")
    ax.legend(frameon=False, loc="upper right", fontsize=13)
    ax.set_title("Pooled scores over 2,048 test pairs", fontsize=15,
                 loc="left", color="#595959")
    fig.tight_layout()
    fig.savefig(path, facecolor="white", bbox_inches="tight")
    plt.close(fig)
    return path


def make_city_chart(rows: list[dict]) -> Path:
    path = ASSETS / "city_f1_slide.png"
    per: dict[str, dict[str, float]] = defaultdict(dict)
    for row in rows:
        cfg = row["configuration"]
        if cfg in ("hardcoded", "percentile"):
            per[row["city"]][cfg] = row["metrics"]["f1"]
    assert len(per) == 10
    cities = sorted(per, key=lambda c: per[c]["percentile"])
    y = list(range(len(cities)))
    fig, ax = plt.subplots(figsize=(8.6, 5.2), dpi=200)
    hard = [per[c]["hardcoded"] for c in cities]
    pct = [per[c]["percentile"] for c in cities]
    ax.scatter(hard, [v - 0.13 for v in y], color="#C55A11", s=95,
               label="Hardcoded NDVI/NDBI", zorder=3)
    ax.scatter(pct, [v + 0.13 for v in y], color="#2E75B6", s=95,
               label="Percentile 99.0", zorder=3)
    ax.axvline(0.0517, color="#7F7F7F", linestyle=":", linewidth=2,
               label="Random at 5.17%")
    ax.set_yticks(y)
    ax.set_yticklabels([c.replace("_", " ") for c in cities], fontsize=13)
    ax.set_xlim(0, 0.35)
    ax.set_xlabel("F1", fontsize=15)
    ax.tick_params(axis="x", labelsize=13)
    ax.grid(axis="x", color="#E6E6E6", linewidth=1)
    ax.set_axisbelow(True)
    for spine in ("top", "right"):
        ax.spines[spine].set_visible(False)
    ax.legend(frameon=False, loc="lower right", fontsize=11)
    ax.set_title("Per-city F1 on the ten OSCD test cities", fontsize=15,
                 loc="left", color="#595959")
    fig.tight_layout()
    fig.savefig(path, facecolor="white", bbox_inches="tight")
    plt.close(fig)
    return path


def make_empty_chart(pooled: dict[int, dict[str, float]]) -> Path:
    path = ASSETS / "empty_mask_slide.png"
    predicted = pooled[20]["predicted"] / (2048 * 256 * 256) * 100
    changed = pooled[20]["changed"] / (2048 * 256 * 256) * 100
    assert round(predicted, 3) == 0.068
    assert round(changed, 2) == 5.09
    fig, ax = plt.subplots(figsize=(6.4, 3.6), dpi=200)
    labels = ["True changed", "Predicted changed"]
    values = [changed, predicted]
    colors = ["#2E75B6", "#C55A11"]
    bars = ax.barh(labels, values, color=colors, height=0.52)
    ax.set_xlim(0, 6.2)
    ax.set_xlabel("Share of pixels at 10 m (%)", fontsize=13)
    ax.tick_params(axis="both", labelsize=13)
    ax.grid(axis="x", color="#E6E6E6", linewidth=1)
    ax.set_axisbelow(True)
    for spine in ("top", "right"):
        ax.spines[spine].set_visible(False)
    for bar, value in zip(bars, values):
        label = f"{value:.3f}%" if value < 0.1 else f"{value:.2f}%"
        ax.text(value + 0.12, bar.get_y() + bar.get_height() / 2,
                label, va="center", fontsize=14, fontweight="bold",
                color="#262626")
    ax.set_title("0.068% predicted vs 5.09% true", fontsize=14,
                 loc="left", color="#595959")
    fig.tight_layout()
    fig.savefig(path, facecolor="white", bbox_inches="tight")
    plt.close(fig)
    return path


def add_text(slide, left, top, width, height, text, *, size=18, color=INK,
             bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             line_spacing=1.0, font=FONT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width),
                                   Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_bullets(slide, left, top, width, height, items, *, size=18,
                color=INK, space_after=12, accent=BLUE):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width),
                                   Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for index, item in enumerate(items):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.space_after = Pt(space_after)
        p.line_spacing = 1.05
        bullet = p.add_run()
        bullet.text = "•  "
        bullet.font.name = FONT
        bullet.font.size = Pt(size)
        bullet.font.bold = True
        bullet.font.color.rgb = accent
        run = p.add_run()
        run.text = item
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.color.rgb = color
    return box


def add_card(slide, left, top, width, height, fill, line=None, radius=0.06):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top),
        Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    shape.adjustments[0] = radius
    shape.shadow.inherit = False
    return shape


def add_title(slide, kicker, title, *, title_size=30):
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(0.42),
        Inches(0.08), Inches(0.86)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = ORANGE
    bar.line.fill.background()
    bar.shadow.inherit = False
    add_text(slide, 0.8, 0.36, 11.8, 0.3, kicker.upper(), size=11,
             color=ORANGE, bold=True)
    add_text(slide, 0.8, 0.65, 11.9, 0.9, title, size=title_size,
             color=NAVY, bold=True, line_spacing=1.0)


def add_footer(slide, number):
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(7.04),
        Inches(12.23), Pt(0.75)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = LINE
    line.line.fill.background()
    line.shadow.inherit = False
    add_text(slide, 0.55, 7.12, 9.5, 0.22,
             "Resolution mismatch case study  ·  saved evidence only",
             size=9, color=MID_GREY)
    add_text(slide, 11.9, 7.12, 0.9, 0.22, f"{number} / 12", size=9,
             color=MID_GREY, align=PP_ALIGN.RIGHT)


def add_picture_fit(slide, path, left, top, max_w, max_h):
    width_px, height_px = Image.open(path).size
    scale = min(max_w / width_px, max_h / height_px)
    width = Inches(width_px * scale)
    height = Inches(height_px * scale)
    x = Inches(left) + int((Inches(max_w) - width) / 2)
    y = Inches(top) + int((Inches(max_h) - height) / 2)
    return slide.shapes.add_picture(str(path), x, y, width=width,
                                    height=height)


def style_cell(cell, text, *, size=13, bold=False, color=INK, fill=WHITE,
               align=PP_ALIGN.LEFT):
    cell.text = ""
    tf = cell.text_frame
    tf.word_wrap = True
    cell.margin_left = Inches(0.07)
    cell.margin_right = Inches(0.06)
    cell.margin_top = Inches(0.03)
    cell.margin_bottom = Inches(0.03)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill


def build_deck() -> tuple[Path, dict[int, str]]:
    pooled = load_degradation()
    evidence, rows = load_oscd()
    ASSETS.mkdir(parents=True, exist_ok=True)
    f1_chart = make_f1_chart(pooled)
    city_chart = make_city_chart(rows)
    empty_chart = make_empty_chart(pooled)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    notes: dict[int, str] = {}

    # ---------------------------------------------------------------- slide 1
    slide = prs.slides.add_slide(blank)
    strip = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), prs.slide_height
    )
    strip.fill.solid()
    strip.fill.fore_color.rgb = NAVY
    strip.line.fill.background()
    strip.shadow.inherit = False
    add_text(slide, 0.8, 0.7, 7.6, 0.35, "EMPIRICAL CASE STUDY  ·  10-MINUTE TALK",
             size=12, color=ORANGE, bold=True)
    add_text(slide, 0.8, 1.25, 7.4, 2.2,
             "Resolution-Induced Performance Collapse of a CNN Change Detector",
             size=37, color=NAVY, bold=True, line_spacing=0.95)
    add_text(slide, 0.8, 3.65, 7.1, 1.0,
             "Simulated and real Sentinel-2 evidence from one frozen checkpoint",
             size=20, color=GREY, line_spacing=1.05)
    add_text(slide, 0.8, 4.95, 6.2, 0.4, "Aravind", size=17, color=INK, bold=True)
    add_text(slide, 0.8, 5.35, 6.8, 0.65,
             "Department of Computer Science and Engineering\nPresidency University, Bengaluru, India",
             size=12, color=GREY, line_spacing=1.1)
    add_text(slide, 0.8, 6.55, 7.0, 0.35,
             "Draft talk  ·  figures generated from saved results  ·  citations pending",
             size=10, color=MID_GREY)
    card = add_card(slide, 8.65, 1.45, 4.0, 4.55, LIGHT_BLUE)
    add_text(slide, 9.05, 1.85, 3.2, 0.3, "POOLED F1 (%)", size=12,
             color=BLUE, bold=True)
    add_text(slide, 9.05, 2.15, 3.2, 1.0, "87.06", size=54, color=BLUE, bold=True)
    add_text(slide, 9.05, 3.15, 3.2, 0.8, "→  4.41", size=36, color=ORANGE,
             bold=True)
    add_text(slide, 9.05, 4.05, 3.2, 0.35, "0.5 m  →  2 m simulated", size=13,
             color=INK, bold=True)
    add_text(slide, 9.05, 4.6, 3.1, 0.9,
             "2,048 LEVIR-CD test pairs\nOne frozen checkpoint",
             size=11, color=GREY, line_spacing=1.2)
    add_footer(slide, 1)
    notes[1] = (
        "Good morning. My talk is about a failure mode that aggregate benchmark "
        "scores hide: a change detector that looks strong at half a metre and "
        "almost stops detecting by two metres. This is a case study of one "
        "frozen checkpoint, not a new model, and not a claim that all CNNs "
        "fail. I will show a controlled degradation experiment on 2,048 aerial "
        "test pairs, a transfer evaluation on ten Sentinel-2 cities, and the "
        "limits of both. The question is narrow: what happens to a validated "
        "score when the input resolution changes? Please hold two questions: "
        "what is actually validated here, and what evidence would change the "
        "conclusion?"
    )

    # ---------------------------------------------------------------- slide 2
    slide = prs.slides.add_slide(blank)
    add_title(slide, "The problem", "A native-resolution score does not travel to 10 m")
    add_bullets(
        slide, 0.8, 1.75, 6.5, 4.4,
        [
            "The checkpoint is trained and scored at 0.5 m on LEVIR-CD aerial pairs.",
            "Sentinel-2 provides 10 m and 20 m bands, free and frequently revisited.",
            "Same model, coarser pixel support: small objects, boundaries and mixed pixels change.",
            "Question: what happens to this checkpoint when the same scenes are coarsened?",
        ],
        size=19, space_after=18,
    )
    panel = add_card(slide, 7.65, 1.8, 5.05, 4.65, LIGHT_BLUE)
    add_text(slide, 8.0, 2.0, 4.3, 0.3, "PIXEL FOOTPRINT, TO SCALE", size=11,
             color=BLUE, bold=True)
    grid_x, grid_y, grid_side = 8.05, 2.45, 2.2
    cell = grid_side / 20
    square = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(grid_x), Inches(grid_y),
        Inches(grid_side), Inches(grid_side)
    )
    square.fill.solid()
    square.fill.fore_color.rgb = PALE_BLUE
    square.line.color.rgb = NAVY
    square.line.width = Pt(1.2)
    square.shadow.inherit = False
    highlight = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(grid_x + 10 * cell),
        Inches(grid_y + 10 * cell), Inches(cell), Inches(cell)
    )
    highlight.fill.solid()
    highlight.fill.fore_color.rgb = ORANGE
    highlight.line.fill.background()
    highlight.shadow.inherit = False
    add_text(slide, 10.5, 2.65, 2.0, 0.75, "10 m\nSentinel-2 pixel", size=14,
             color=INK, bold=True, line_spacing=1.1)
    add_text(slide, 10.5, 3.6, 1.95, 0.8,
             "0.5 m LEVIR-CD pixel\n(orange square)", size=11, color=GREY,
             line_spacing=1.1)
    add_text(slide, 8.0, 4.85, 4.3, 0.4, "20× coarser linearly", size=20,
             color=ORANGE, bold=True)
    add_text(slide, 8.0, 5.22, 4.3, 0.3, "400× the area per pixel", size=14,
             color=NAVY, bold=True)
    add_text(slide, 8.0, 5.58, 4.25, 0.75,
             "The orange square is one 0.5 m LEVIR-CD pixel; the full square is "
             "one 10 m Sentinel-2 pixel.", size=11, color=GREY,
             line_spacing=1.15)
    add_footer(slide, 2)
    notes[2] = (
        "The starting point is a mismatch of pixel support. The checkpoint was "
        "trained and scored on LEVIR-CD, which is half-metre aerial imagery. "
        "Sentinel-2 gives us ten-metre and twenty-metre bands, free and "
        "frequently revisited. Ten metres is twenty times coarser per side than "
        "half a metre, and four hundred times the area. On this diagram, the "
        "large square is one Sentinel-2 pixel; the orange square is one "
        "LEVIR-CD pixel, drawn to scale. At that size, small buildings, road "
        "edges and boundaries become sub-pixel or mixed. That is not a cosmetic "
        "difference. The question is whether a model validated on the fine grid "
        "still works when the input is coarsened. The paper does not assume the "
        "answer; it measures the same checkpoint under a controlled transform."
    )

    # ---------------------------------------------------------------- slide 3
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Design", "One frozen model, two evaluations, no retraining")
    left = add_card(slide, 0.65, 1.7, 5.95, 4.2, LIGHT_BLUE)
    add_text(slide, 1.0, 1.95, 5.2, 0.4, "Controlled degradation", size=21,
             color=BLUE, bold=True)
    add_bullets(
        slide, 1.0, 2.5, 5.2, 3.1,
        [
            "2,048 LEVIR-CD test pairs",
            "Six factors: 1, 2, 4, 8, 16, 20",
            "Gaussian blur + LANCZOS down/up",
            "Same scene content; only sampling changes",
        ],
        size=16, space_after=10, accent=BLUE,
    )
    right = add_card(slide, 6.9, 1.7, 5.8, 4.2, LIGHT_ORANGE)
    add_text(slide, 7.25, 1.95, 5.1, 0.4, "Real Sentinel-2 (OSCD)", size=21,
             color=ORANGE, bold=True)
    add_bullets(
        slide, 7.25, 2.5, 5.1, 3.1,
        [
            "10 official test cities",
            "3,077,936 pixels; 5.17% changed",
            "Zero-shot CNN; NDVI/NDBI and Otsu rules",
            "No SCL; confounded with sensor and season",
        ],
        size=16, space_after=10, accent=ORANGE,
    )
    strip2 = add_card(slide, 0.65, 6.15, 12.05, 0.55, PALE_BLUE)
    add_text(slide, 0.95, 6.27, 11.5, 0.3,
             "Frozen checkpoint  ·  one architecture  ·  one training run  ·  fixed 0.5 threshold",
             size=13, color=NAVY, bold=True)
    add_footer(slide, 3)
    notes[3] = (
        "Here is the design. Two evaluations, one frozen checkpoint, no "
        "retraining. On the left, the 2,048 LEVIR-CD test pairs at six factors "
        "from one to twenty. On the right, the ten official OSCD test cities: "
        "3,077,936 pixels, 5.17 percent labelled change. The CNN weights are "
        "zero-shot on OSCD; the spectral rules are train-selected reference "
        "points, not a proposed solution. One architecture, one training run, "
        "threshold fixed at 0.5 throughout. That scope is deliberate. This is a "
        "case study of a deployed checkpoint, not an architecture comparison, "
        "and I will not generalise from one model to all CNNs."
    )

    # ---------------------------------------------------------------- slide 4
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Degradation protocol", "Coarsen, then upsample back to 256×256")
    steps = [
        ("1", "256×256", "before / after pair"),
        ("2", "Gaussian blur", "radius f / 2"),
        ("3", "LANCZOS down", "to (256/f)²"),
        ("4", "LANCZOS up", "back to 256×256"),
    ]
    fills = [NAVY, BLUE, TEAL, ORANGE]
    x = 0.65
    for index, (number, head, sub) in enumerate(steps):
        box = add_card(slide, x, 1.85, 2.68, 1.3, fills[index], radius=0.08)
        add_text(slide, x + 0.18, 2.0, 0.35, 0.3, number, size=13, color=WHITE,
                 bold=True)
        add_text(slide, x + 0.18, 2.28, 2.3, 0.35, head, size=16, color=WHITE,
                 bold=True)
        add_text(slide, x + 0.18, 2.68, 2.3, 0.3, sub, size=11,
                 color=RGBColor(0xEA, 0xEA, 0xEA))
        if index < 3:
            arrow = slide.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW, Inches(x + 2.72), Inches(2.35),
                Inches(0.26), Inches(0.3)
            )
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = ORANGE
            arrow.line.fill.background()
            arrow.shadow.inherit = False
        x += 3.0
    band = add_card(slide, 0.65, 3.55, 12.05, 0.6, PALE_BLUE)
    add_text(slide, 0.95, 3.69, 11.5, 0.32,
             "Factors 1, 2, 4, 8, 16, 20   →   0.5, 1, 2, 4, 8, 10 m",
             size=17, color=NAVY, bold=True)
    add_card(slide, 0.65, 4.45, 5.75, 0.95, LIGHT_TEAL)
    add_text(slide, 0.95, 4.62, 5.2, 0.3, "At 10 m: a 12×12 intermediate", size=17,
             color=TEAL, bold=True)
    add_text(slide, 0.95, 4.97, 5.2, 0.28, "for a 256×256 input", size=12, color=GREY)
    add_card(slide, 6.65, 4.45, 6.05, 0.95, LIGHT_BLUE)
    add_text(slide, 6.95, 4.62, 5.5, 0.3, "ImageNet normalization · threshold 0.5",
             size=17, color=BLUE, bold=True)
    add_text(slide, 6.95, 4.97, 5.5, 0.28, "pooled TP / FP / FN over 2,048 pairs",
             size=12, color=GREY)
    warn = add_card(slide, 0.65, 5.75, 12.05, 0.85, LIGHT_ORANGE)
    add_text(slide, 0.95, 5.95, 11.5, 0.45,
             "This isolates blur and resampling. It is not a Sentinel-2 simulation.",
             size=18, color=ORANGE, bold=True)
    add_footer(slide, 4)
    notes[4] = (
        "The degradation protocol is simple and reproducible. Start with a 256 "
        "by 256 pair. Blur with a Gaussian whose radius is the factor divided "
        "by two. Downsample with LANCZOS to 256 over the factor, squared. "
        "Upsample back to 256. At factor twenty, the intermediate is only "
        "twelve by twelve pixels. Factors one, two, four, eight, sixteen and "
        "twenty map to half a metre through ten metres. Then ImageNet "
        "normalization, a fixed 0.5 threshold, and pooled true positives, false "
        "positives and false negatives over all 2,048 pairs. The blur and "
        "resampling isolate the effect of pixel support. They do not simulate "
        "Sentinel-2 spectra, atmosphere, mixed-pixel physics or noise. I will "
        "come back to that distinction."
    )

    # ---------------------------------------------------------------- slide 5
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Result", "F1 falls from 87.06% to 4.41% by 2 m")
    add_picture_fit(slide, f1_chart, 0.55, 1.7, 7.9, 4.9)
    panel5 = add_card(slide, 8.75, 1.75, 3.95, 4.85, LIGHT_BLUE)
    add_text(slide, 9.05, 2.0, 3.3, 0.3, "POOLED F1 (%)", size=11, color=BLUE,
             bold=True)
    add_text(slide, 9.05, 2.3, 3.3, 0.55, "87.06  →  4.41", size=25, color=NAVY,
             bold=True)
    rows5 = [
        ("0.5 m", "87.06"),
        ("1 m", "70.81  (−16.25)"),
        ("2 m", "4.41"),
        ("8–10 m", "2.61"),
    ]
    y = 3.05
    for label, value in rows5:
        add_text(slide, 9.05, y, 1.1, 0.3, label, size=13, color=GREY, bold=True)
        add_text(slide, 10.2, y, 2.3, 0.3, value, size=13, color=INK)
        y += 0.42
    add_text(slide, 9.05, 4.95, 3.3, 0.3, "Recall: 85.36% → 1.32%", size=13,
             color=ORANGE, bold=True)
    add_text(slide, 9.05, 5.35, 3.3, 0.85,
             "The decline is recall-dominated.\nPooled counts; one checkpoint.",
             size=11, color=GREY, line_spacing=1.15)
    add_footer(slide, 5)
    notes[5] = (
        "Here is the main result. At half a metre, pooled F1 is 87.06 percent. "
        "At one metre it is 70.81, a drop of 16.25 points. At two metres it is "
        "4.41 percent. From eight to ten metres it is 2.61. The dashed line is "
        "recall: 85.36 percent at native resolution, 57.01 at one metre, 2.26 "
        "at two metres, and 1.32 at eight to ten. The decline is "
        "recall-dominated. Precision actually rises, from 88.8 percent at "
        "native resolution to 99.5 percent at ten metres. So the model is not "
        "becoming noisy; it is becoming silent. That is why the next slide "
        "matters more than the F1 number alone."
    )

    # ---------------------------------------------------------------- slide 6
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Why precision misleads", "At 10 m, the model predicts almost nothing")
    stats = [
        ("0.068%", "of pixels predicted changed at 10 m\n91,006 of 134,217,728", ORANGE),
        ("924 / 935", "changed images with zero predicted change\n(98.82%)", BLUE),
        ("99.47%", "precision, from 90,527 TP and 479 FP\nthe mask is nearly empty", TEAL),
    ]
    y = 1.72
    for value, label, color in stats:
        add_card(slide, 0.65, y, 5.65, 1.35, WHITE, line=LINE)
        add_text(slide, 0.95, y + 0.15, 2.2, 0.55, value, size=27, color=color,
                 bold=True)
        add_text(slide, 3.15, y + 0.18, 2.9, 0.95, label, size=11, color=GREY,
                 line_spacing=1.15)
        y += 1.55
    add_picture_fit(slide, empty_chart, 6.65, 1.7, 6.05, 4.7)
    add_text(slide, 0.65, 6.45, 12.05, 0.4,
             "High precision means the mask is nearly empty — not that localisation is good.",
             size=18, color=NAVY, bold=True)
    add_footer(slide, 6)
    notes[6] = (
        "At ten metres the model predicts 91,006 pixels changed out of "
        "134,217,728. That is 0.068 percent. 924 of the 935 images that "
        "actually contain change receive no predicted change at all, which is "
        "98.82 percent. Precision is 99.47 percent, but it comes from 90,527 "
        "true positives and only 479 false positives. A precision-only summary "
        "would call this excellent; a recall or predicted-fraction summary "
        "exposes that the mask is almost empty. This is the central reporting "
        "lesson: precision is not interpretable without prevalence and "
        "predicted fraction. The paper therefore reports all three."
    )

    # ---------------------------------------------------------------- slide 7
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Real Sentinel-2", "OSCD transfer: near-empty CNN, weak spectral rules",
             title_size=28)
    table_shape = slide.shapes.add_table(7, 6, Inches(0.6), Inches(1.65),
                                         Inches(12.15), Inches(3.95))
    table = table_shape.table
    widths = [3.05, 1.25, 1.3, 1.15, 2.55, 2.85]
    for index, width in enumerate(widths):
        table.columns[index].width = Inches(width)
    table.rows[0].height = Inches(0.45)
    for index in range(1, 7):
        table.rows[index].height = Inches(0.5)
    headers = ["Method", "Pred. %", "Precision", "Recall", "Score", "95% CI (F1)"]
    for col, header in enumerate(headers):
        style_cell(table.cell(0, col), header, size=12, bold=True, color=WHITE,
                   fill=NAVY)
    data = [
        ("Hardcoded NDVI/NDBI", "2.24", "0.1292", "0.0561", "F1 0.0782", "0.0168–0.1846"),
        ("Percentile 99.0", "1.66", "0.2419", "0.0777", "F1 0.1176", "0.0540–0.1802"),
        ("Otsu", "67.35", "0.0513", "0.6684", "OA 0.3439", "F1 0.0953; 0.0495–0.1245"),
        ("CNN joint percentile", "0.01", "0.2983", "0.0003", "F1 0.0007", "0.0000–0.0017"),
        ("CNN reflectance", "0.00", "—", "0.0000", "F1 0.0000", "—"),
        ("Random at prevalence", "5.17", "0.0517", "0.0517", "F1 0.0517", "—"),
    ]
    for r, row_data in enumerate(data, start=1):
        fill = LIGHT_ORANGE if row_data[0] == "Otsu" else (WHITE if r % 2 else LIGHT_BLUE)
        for c, value in enumerate(row_data):
            bold = c == 0 or (c == 4 and r <= 3)
            style_cell(table.cell(r, c), value, size=11.5, bold=bold,
                       color=INK, fill=fill)
    add_text(slide, 0.65, 5.85, 12.1, 0.35,
             "True change 5.17%  ·  10 official test cities  ·  3,077,936 pixels",
             size=15, color=NAVY, bold=True)
    add_text(slide, 0.65, 6.28, 12.1, 0.6,
             "The test cities had been scored before the reported run. These are descriptive scores, not an independent test.",
             size=13, color=ORANGE, bold=True)
    add_footer(slide, 7)
    notes[7] = (
        "Now the real Sentinel-2 evaluation. OSCD, ten official test cities, "
        "3,077,936 pixels, 5.17 percent true change. The percentile rule has "
        "the best pooled F1 among the spectral rules, 0.1176, with a "
        "city-bootstrap interval from 0.054 to 0.180. The hardcoded NDVI/NDBI "
        "rule is 0.0782, interval 0.017 to 0.185. Those intervals overlap, and "
        "the paired difference interval includes zero, so I do not rank them. "
        "Otsu predicts 67.35 percent of pixels changed; its headline is overall "
        "accuracy 0.3439, not F1. The CNN is near empty: F1 0.0007 for "
        "joint-percentile preprocessing, which was chosen on the training "
        "cities, and 0.0000 for reflectance. One disclosure is essential: these "
        "ten cities had been scored before the reported run. These are "
        "descriptive scores, not an independent test, and the paper says so."
    )

    # ---------------------------------------------------------------- slide 8
    slide = prs.slides.add_slide(blank)
    add_title(slide, "City-level variation", "Pooled scores hide large, inconsistent city differences",
             title_size=28)
    add_picture_fit(slide, city_chart, 0.5, 1.65, 7.8, 4.9)
    panel8 = add_card(slide, 8.55, 1.75, 4.15, 4.8, LIGHT_BLUE)
    add_text(slide, 8.85, 1.98, 3.55, 0.3, "F1 ACROSS 10 CITIES", size=11,
             color=BLUE, bold=True)
    items = [
        ("Percentile", "0.0005 – 0.2481"),
        ("Hardcoded", "0.0032 – 0.3102"),
        ("Random (analytic)", "0.0517"),
        ("CNN joint", "zero in 8 / 10 cities"),
    ]
    y = 2.4
    for label, value in items:
        add_text(slide, 8.85, y, 3.6, 0.28, label, size=12, color=GREY, bold=True)
        add_text(slide, 8.85, y + 0.27, 3.6, 0.35, value, size=16, color=NAVY,
                 bold=True)
        y += 0.85
    add_text(slide, 8.85, 5.75, 3.55, 0.65,
             "Descriptive ranges; ten cities. Intervals resample cities or pixels, not training runs.",
             size=10.5, color=GREY, line_spacing=1.15)
    add_footer(slide, 8)
    notes[8] = (
        "Pooled numbers hide the city variation. Percentile F1 ranges from "
        "0.0005 in saclay_w to 0.2481 in chongqing. Hardcoded ranges from "
        "0.0032 in dubai to 0.3102 in chongqing. The random-at-prevalence F1 is "
        "0.0517, shown as the dotted line; many city values sit near or below "
        "it. The CNN predicts nothing in eight of ten cities. These are "
        "descriptive ranges from ten cities. The intervals resample cities or "
        "pixels; neither includes training or threshold-selection uncertainty. "
        "So the pooled score is a summary, not a stable estimate of how the "
        "tool will behave in a new city."
    )

    # ---------------------------------------------------------------- slide 9
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Interpretation", "Test the resolution path you will deploy")
    points = [
        ("1", "Test the deployed resolution path.",
         "A native-resolution benchmark score does not reveal coarse-resolution failure."),
        ("2", "Report precision with predicted fraction and prevalence.",
         "99.47% precision meant 0.068% of pixels predicted changed."),
        ("3", "Do not treat simple indices as a validated fallback.",
         "Both spectral rules under-predict; Otsu predicts 67.35% changed."),
    ]
    y = 1.85
    for number, head, sub in points:
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(0.8), Inches(y), Inches(0.64), Inches(0.64)
        )
        circle.fill.solid()
        circle.fill.fore_color.rgb = ORANGE
        circle.line.fill.background()
        circle.shadow.inherit = False
        add_text(slide, 0.8, y + 0.14, 0.64, 0.3, number, size=17, color=WHITE,
                 bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, 1.75, y - 0.02, 10.7, 0.4, head, size=20, color=NAVY,
                 bold=True)
        add_text(slide, 1.75, y + 0.4, 10.7, 0.35, sub, size=15, color=GREY)
        y += 1.45
    bottom = add_card(slide, 0.8, 6.15, 11.9, 0.55, PALE_BLUE)
    add_text(slide, 1.05, 6.27, 11.4, 0.3,
             "The paper identifies a risk in evaluation, not a new detector.",
             size=16, color=NAVY, bold=True)
    add_footer(slide, 9)
    notes[9] = (
        "What does this mean in practice? Start with the resolution path you "
        "will deploy. A native-resolution benchmark score does not reveal a "
        "coarse-resolution failure; if the product runs on ten-metre data, the "
        "evaluation should too. Second, report precision with predicted "
        "fraction and prevalence. A near-empty mask can look nearly perfect. "
        "Third, do not treat simple spectral indices as a validated fallback. "
        "The percentile rule is the best of the spectral options here, but it "
        "under-predicts change, and Otsu predicts two thirds of pixels as "
        "changed. The paper identifies an evaluation risk. It does not explain "
        "the real Sentinel-2 gap or propose a new detector."
    )

    # --------------------------------------------------------------- slide 10
    slide = prs.slides.add_slide(blank)
    add_title(slide, "What this does not show", "One checkpoint, one synthetic degradation, a used test split",
             title_size=28)
    add_bullets(
        slide, 0.8, 1.8, 7.3, 4.4,
        [
            "One architecture and one training run; no retraining or seed variation.",
            "Synthetic degradation: no spectra, atmosphere, mixed pixels or noise.",
            "OSCD confounds resolution with sensor, geography and season; no SCL.",
            "Ten cities; the test split had been scored before the reported run.",
            "The deployed app and scene selection are not validated end to end.",
        ],
        size=17, space_after=16, accent=ORANGE,
    )
    caution = add_card(slide, 8.4, 1.85, 4.3, 3.3, LIGHT_ORANGE)
    add_text(slide, 8.75, 2.1, 3.6, 0.35, "CLAIM BOUNDARY", size=11, color=ORANGE,
             bold=True)
    add_text(slide, 8.75, 2.55, 3.6, 2.2,
             "No causal claim\nabout the Sentinel-2 gap\n\nNo priority claim\n\nNo operational claim",
             size=18, color=NAVY, bold=True, line_spacing=1.0)
    add_text(slide, 8.75, 5.35, 3.6, 0.8,
             "The paper reports descriptive evidence, not a validated system.",
             size=11, color=GREY, line_spacing=1.15)
    add_footer(slide, 10)
    notes[10] = (
        "The limits matter. One architecture and one training run; no "
        "retraining, fine-tuning or seed variation. The degradation is "
        "synthetic: no Sentinel-2 spectra, atmosphere, mixed-pixel physics or "
        "noise. The OSCD evaluation uses real data, but it confounds resolution "
        "with sensor, geography and season, and the benchmark layout has no "
        "SCL, so clouds are uncontrolled. Ten cities, and the test split had "
        "been scored before the reported run. Finally, the deployed "
        "application, its scene selection and its interface are not validated "
        "end to end. So: no causal claim about the real Sentinel-2 gap, no "
        "priority claim, no operational claim. What remains is descriptive "
        "evidence of an evaluation risk."
    )

    # --------------------------------------------------------------- slide 11
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Next evidence", "What would make the result stronger")
    future = [
        ("1", "A second architecture and checkpoint, with multiple seeds."),
        ("2", "Real same-scene acquisitions at multiple resolutions."),
        ("3", "An untouched city set or a new dataset."),
        ("4", "An SCL-preserving benchmark and a wider threshold grid."),
    ]
    positions = [(0.7, 1.8), (6.8, 1.8), (0.7, 3.55), (6.8, 3.55)]
    for (number, text), (x, y) in zip(future, positions):
        add_card(slide, x, y, 5.8, 1.5, LIGHT_BLUE)
        add_text(slide, x + 0.25, y + 0.2, 0.45, 0.35, number, size=17,
                 color=BLUE, bold=True)
        add_text(slide, x + 0.8, y + 0.2, 4.7, 1.1, text, size=17, color=NAVY,
                 bold=True, line_spacing=1.05)
    add_text(slide, 0.7, 5.55, 11.9, 0.35,
             "An operator study on whether warnings change decisions.",
             size=15, color=GREY)
    add_text(slide, 0.7, 6.15, 11.9, 0.4,
             "These are evidence needs, not results.",
             size=20, color=ORANGE, bold=True)
    add_footer(slide, 11)
    notes[11] = (
        "What would change the conclusion? A second architecture and "
        "checkpoint, with multiple seeds. Real acquisitions of the same scenes "
        "at multiple resolutions, rather than synthetic resampling. An "
        "untouched city set or a new dataset. An SCL-preserving benchmark and a "
        "wider threshold grid. And an operator study on whether warnings "
        "actually change decisions. These are evidence needs, not results; "
        "none of them is in the paper. If you take one thing for your own "
        "evaluation practice, prioritise these two: test the deployed path, "
        "and do not reuse a test split."
    )

    # --------------------------------------------------------------- slide 12
    slide = prs.slides.add_slide(blank)
    add_text(slide, 0.8, 0.75, 11.8, 0.35, "TAKEAWAY", size=12, color=ORANGE,
             bold=True)
    add_text(slide, 0.8, 1.25, 11.8, 1.2,
             "A 0.5 m score is not a 10 m validation.",
             size=38, color=NAVY, bold=True, line_spacing=0.95)
    stats12 = [
        ("87.06% → 4.41%", "pooled F1 from 0.5 m to 2 m\nunder blur + resampling", ORANGE),
        ("0.068%", "of pixels predicted changed at 10 m\n91,006 of 134,217,728", BLUE),
        ("924 / 935", "changed images with zero\npredicted change (98.82%)", TEAL),
    ]
    x = 0.7
    for value, label, color in stats12:
        add_card(slide, x, 2.85, 3.85, 1.75, WHITE, line=LINE)
        add_text(slide, x + 0.25, 3.02, 3.35, 0.55, value, size=22, color=color,
                 bold=True)
        add_text(slide, x + 0.25, 3.65, 3.35, 0.8, label, size=11, color=GREY,
                 line_spacing=1.15)
        x += 4.05
    add_text(slide, 0.8, 5.15, 11.8, 0.4,
             "Real OSCD: CNN F1 0.0007  ·  spectral rules 0.078–0.118  ·  neither is a validated fallback.",
             size=17, color=NAVY, bold=True)
    add_text(slide, 0.8, 5.85, 11.8, 0.35,
             "Paper, figures and artifact: TODO — insert verified repository URL and DOI.",
             size=13, color=GREY)
    add_text(slide, 0.8, 6.25, 11.8, 0.35,
             "Draft; citations and PDF pending.  Thank you.",
             size=13, color=MID_GREY)
    add_footer(slide, 12)
    notes[12] = (
        "Takeaway: a half-metre score is not a ten-metre validation. F1 fell "
        "from 87.06 to 4.41 percent by two metres under blur and resampling. At "
        "ten metres, 0.068 percent of pixels were predicted changed, and 924 of "
        "935 changed images were empty. On real OSCD, the CNN F1 was 0.0007, "
        "the spectral rules were 0.078 to 0.118, and neither is a validated "
        "fallback. The paper reports these limits explicitly. The evidence is "
        "saved, the code and figures are in the artifact, and citations are "
        "still being verified. Thank you; I am happy to take questions."
    )

    assert len(prs.slides) == 12
    for index, slide in enumerate(prs.slides, start=1):
        slide.notes_slide.notes_text_frame.text = notes[index]

    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / "slides.pptx"
    prs.save(path)
    return path, notes


def check_layout(path: Path) -> None:
    prs = Presentation(path)
    assert len(prs.slides) == 12
    width, height = prs.slide_width, prs.slide_height
    for index, slide in enumerate(prs.slides, start=1):
        assert slide.has_notes_slide
        assert slide.notes_slide.notes_text_frame.text.strip()
        for shape in slide.shapes:
            if shape.left is None:
                continue
            if shape.left < 0 or shape.top < 0:
                raise AssertionError(f"slide {index}: shape starts outside canvas")
            if shape.left + shape.width > width + 10000:
                raise AssertionError(f"slide {index}: shape exceeds width")
            if shape.top + shape.height > height + 10000:
                raise AssertionError(f"slide {index}: shape exceeds height")
    print(f"slides={len(prs.slides)} size={width}x{height} layout_ok=True")


def write_speaker_notes(notes: dict[int, str]) -> Path:
    topics = [
        "Title and roadmap",
        "Why resolution is the problem",
        "Design: one checkpoint, two evaluations",
        "Degradation protocol",
        "F1 versus resolution",
        "Why precision misleads",
        "Real Sentinel-2 (OSCD) results",
        "City-level variation",
        "Practical interpretation",
        "What this does not show",
        "What evidence would strengthen it",
        "Takeaway",
    ]
    targets = [25, 50, 45, 55, 55, 50, 70, 50, 55, 50, 40, 40]
    words = {index: len(notes[index].split()) for index in range(1, 13)}
    total_words = sum(words.values())
    path = OUT / "speaker_notes.md"
    lines = [
        "# Speaker notes - 10-minute talk",
        "",
        "Generated from the notes embedded in `slides.pptx`. Edit the generator,",
        "not this file, so the deck and script stay in sync.",
        "",
        f"Speaking budget: **9:45** plus **0:15** buffer. The script is "
        f"**{total_words} words**, about **{total_words / 130:.1f} minutes** at "
        "130 words per minute. Rehearse with a timer; do not speed up to fit.",
        "",
        "## Timing",
        "",
        "| Slide | Topic | Target | Words | Cumulative |",
        "|---:|---|---:|---:|---:|",
    ]
    cumulative = 0
    for index, (topic, target) in enumerate(zip(topics, targets), start=1):
        cumulative += target
        lines.append(
            f"| {index} | {topic} | {target // 60}:{target % 60:02d} | "
            f"{words[index]} | {cumulative // 60}:{cumulative % 60:02d} |"
        )
    lines += [
        "",
        "If you are behind after slide 6, compress slides 8 and 11 to one sentence",
        "each. Never cut the repeated-test disclosure on slide 7 or the claim",
        "boundary on slide 10. If you are ahead, stop; do not add results.",
        "",
        "## Script",
        "",
    ]
    for index, topic in enumerate(topics, start=1):
        lines += [
            f"### Slide {index} - {topic} ({targets[index - 1]} s)",
            "",
            notes[index].strip(),
            "",
        ]
    lines += [
        "## Delivery notes",
        "",
        "- Pause after the 4.41% number on slide 5 and after the 0.068% number",
        "  on slide 6. Those are the two results the audience should retain.",
        "- Say the repeated-test disclosure on slide 7 slowly and without",
        "  qualification. Do not bury it in the limitations slide.",
        "- On slide 2, point to the orange square, not the large blue square.",
        "- Do not improvise a published baseline, a v3 score, a fine-tuned",
        "  result, a causal mechanism or an operational claim. If asked, use",
        "  `QA.md` and name the missing artifact.",
        "- The deck says `Draft` because citations are not yet verified and no",
        "  PDF build exists. Do not remove that label for a public talk.",
        "",
        "Q&A: `QA.md`. Demo: `DEMO.md`. Slide-level claim audit:",
        "`CLAIMS_ON_SLIDES.md`. One-page handout: `handout.pdf`.",
        "",
    ]
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


if __name__ == "__main__":
    deck, notes = build_deck()
    check_layout(deck)
    notes_path = write_speaker_notes(notes)
    print(f"wrote {deck}")
    print(f"wrote {notes_path}")
